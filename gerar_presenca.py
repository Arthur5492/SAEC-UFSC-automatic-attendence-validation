"""Gera a presença de cada disciplina a partir dos check-ins da SAEC no Even3.

Uso: python gerar_presenca.py
"""
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent
PLANILHA = next(BASE.glob("ListaParticipante_*.xlsx"))
PASTA_EVENTOS = BASE / "presenca_em_cada_evento"
PASTA_DISCIPLINAS = BASE / "disciplinas"
METADATA = PASTA_DISCIPLINAS / "metadata.json"
RELATORIO = BASE / "relatorio.txt"

INICIO_SAEC = date(2026, 10, 5)  # segunda-feira da semana acadêmica
DURACAO_PERIODO = timedelta(minutes=50)
DIAS = {
    2: "segunda-feira",
    3: "terça-feira",
    4: "quarta-feira",
    5: "quinta-feira",
    6: "sexta-feira",
}

# Na sexta teve hackathon com vagas limitadas, então nem todos puderam participar
DIA_PRESENCA_OPCIONAL = 6

CORES = {"presente": "C6EFCE", "ausente": "FFC7CE", "presença opcional": "E7E6E6"}

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

relatorio = defaultdict(list)


def aviso(secao, msg):
    print(f"AVISO: {msg}")
    relatorio[secao].append(msg)


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(texto.lower().split())


def somente_digitos(texto):
    return re.sub(r"\D", "", str(texto or ""))


#? Inscrições do Even3: nome/e-mail -> matrícula

def carregar_inscricoes():
    linhas = list(openpyxl.load_workbook(PLANILHA, read_only=True).active.iter_rows(values_only=True))
    cab = [normalizar(c) for c in linhas[0]]
    i_nome, i_email, i_mat = cab.index("nome"), cab.index("e-mail"), cab.index("numero da matricula")

    por_nome, por_email = {}, {}
    contagem = defaultdict(int)
    for linha in linhas[1:]:
        nome = str(linha[i_nome] or "").strip()
        if not nome:
            continue
        email = str(linha[i_email] or "").strip().lower()
        mat_bruta = str(linha[i_mat] or "").strip()
        mat = somente_digitos(mat_bruta)
        chave = normalizar(nome)
        contagem[chave] += 1

        # tem gente inscrita duas vezes, junta e fica com a matrícula preenchida
        pessoa = por_nome.get(chave) or por_email.get(email) or {"nome": nome, "matricula": "", "bruta": ""}
        if mat and not pessoa["matricula"]:
            pessoa["matricula"], pessoa["bruta"] = mat, mat_bruta
        por_nome[chave] = pessoa
        if email:
            por_email[email] = pessoa

    for chave, n in contagem.items():
        if n > 1:
            p = por_nome[chave]
            relatorio["Inscrições duplicadas no Even3 (mescladas)"].append(
                f"{p['nome']} ({n}x) -> matrícula {p['matricula'] or 'nenhuma'}")
    return por_nome, por_email


#? Listas de presença (PDF): matrícula -> eventos

def carregar_eventos():
    eventos = []
    for pdf in sorted(PASTA_EVENTOS.glob("*.pdf")):
        with pdfplumber.open(pdf) as doc:
            texto = "\n".join(pg.extract_text() or "" for pg in doc.pages)
        m = re.search(r"(\d{2}/\d{2}/\d{4}) às (\d{2}:\d{2}) - (\d{2}:\d{2})", texto)
        if not m:
            aviso("Erros de leitura", f"período não encontrado em {pdf.name}")
            continue
        dia = datetime.strptime(m.group(1), "%d/%m/%Y")
        inicio = datetime.combine(dia, datetime.strptime(m.group(2), "%H:%M").time())
        fim = datetime.combine(dia, datetime.strptime(m.group(3), "%H:%M").time())
        participantes = []
        for linha in texto.splitlines():
            p = re.match(r"(.+?)\s+(\S+@\S+)$", linha.strip())
            if p:
                participantes.append((p.group(1).strip(), p.group(2).lower()))
        eventos.append({"nome": pdf.stem, "inicio": inicio, "fim": fim, "participantes": participantes})
    return eventos


def mapear_presencas(eventos, por_nome, por_email):
    presencas = defaultdict(list)
    avisados = set()
    for ev in eventos:
        for nome, email in ev["participantes"]:
            pessoa = por_nome.get(normalizar(nome)) or por_email.get(email)
            if pessoa is None:
                if nome not in avisados:
                    aviso("Participantes não encontrados na planilha", f"aluno {nome} ({email}) não está na lista de inscritos")
                    avisados.add(nome)
                continue
            if not pessoa["matricula"]:
                if nome not in avisados:
                    aviso("Participantes sem matrícula (descartados)", f"aluno {nome} não contém matrícula")
                    avisados.add(nome)
                continue
            presencas[pessoa["matricula"]].append(ev)
    return presencas


#? Turmas e horários das aulas

def parse_horario(horario):
    # formato do CAGR: "2.1420-2 / CTS-SL313A" = segunda, 14:20, 2 períodos
    m = re.match(r"\s*(\d)\.(\d{2})(\d{2})-(\d+)", horario)
    dia_semana, hh, mm, periodos = (int(g) for g in m.groups())
    dia = INICIO_SAEC + timedelta(days=dia_semana - 2)
    inicio = datetime(dia.year, dia.month, dia.day, hh, mm)
    return dia_semana, inicio, inicio + periodos * DURACAO_PERIODO


def esteve_na_aula(eventos_aluno, blocos):
    return any(ev["inicio"] < fim and inicio < ev["fim"]
               for ev in eventos_aluno for inicio, fim in blocos)


def ler_alunos(pasta):
    arquivo = next(f for f in pasta.glob("*.csv") if f.name != "presença.csv")
    alunos = []
    with open(arquivo, encoding="utf-8-sig", newline="") as f:
        for linha in csv.DictReader(f):
            # o Moodle exporta a matrícula no campo Sobrenome, entre parênteses
            mat = re.fullmatch(r"\s*\((\d+)\)\s*", linha["Sobrenome"])
            if not mat:  # professores e monitores
                nome = f"{linha['Nome']} {linha['Sobrenome']}".strip()
                aviso("Linhas das turmas sem matrícula (ignoradas)", f"{pasta.name}: {nome} não contém matrícula na lista da turma")
                continue
            alunos.append((linha["Nome"].strip(), mat.group(1), linha["Endereço de email"].strip().lower()))
    return alunos


def corrigir_matriculas(presencas, por_nome, por_email, turmas):
    # Se a matrícula do Even3 não está em nenhuma turma mas o nome ou e-mail bate,
    # a pessoa digitou errado. Usa a matrícula da turma.
    matriculas_em_turmas, turma_por_nome, turma_por_email = set(), {}, {}
    for disc, alunos in turmas:
        for nome, mat, email in alunos:
            matriculas_em_turmas.add(mat)
            turma_por_nome[normalizar(nome)] = (mat, disc["id"])
            turma_por_email[email] = (mat, disc["id"])

    pessoas = {id(p): p for p in por_nome.values()}.values()
    for p in sorted(pessoas, key=lambda p: normalizar(p["nome"])):
        if p["matricula"] not in presencas or p["matricula"] in matriculas_em_turmas:
            continue
        correta = turma_por_nome.get(normalizar(p["nome"])) or next(
            (turma_por_email[e] for e, q in por_email.items() if q is p and e in turma_por_email), None)
        if correta:
            aviso("Matrícula digitada errada no Even3 (corrigida pela lista da turma)",
                  f"{p['nome']} informou {p['bruta']}, mas na turma {correta[1]} consta {correta[0]} - presença contabilizada")
            presencas[correta[0]].extend(presencas.pop(p["matricula"]))
        else:
            relatorio["Participantes com matrícula que não está em nenhuma turma"].append(
                f"{p['nome']} - matrícula informada: {p['bruta']}")


def montar_tabela(disc, alunos, presencas):
    blocos_por_dia = defaultdict(list)
    for h in disc["horarios"]:
        dia_semana, inicio, fim = parse_horario(h)
        blocos_por_dia[dia_semana].append((inicio, fim))
    dias = sorted(blocos_por_dia)

    cabecalho = ["Nome", "Matrícula"] + [DIAS[d] for d in dias]
    linhas, presentes = [], 0
    for nome, mat, _ in alunos:
        linha, foi = [nome, mat], False
        for d in dias:
            if d == DIA_PRESENCA_OPCIONAL:
                linha.append("presença opcional")
            elif esteve_na_aula(presencas.get(mat, []), blocos_por_dia[d]):
                linha.append("presente")
                foi = True
            else:
                linha.append("ausente")
        linhas.append(linha)
        presentes += foi
    return cabecalho, linhas, presentes


#? Saída

def salvar_csv(caminho, cabecalho, linhas):
    with open(caminho, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(cabecalho)
        w.writerows(linhas)


def salvar_xlsx(caminho, cabecalho, linhas, disc):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = disc["id"]
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append(linha)
        for c in ws[ws.max_row][2:]:
            c.fill = PatternFill("solid", fgColor=CORES[c.value])
    ws.freeze_panes = "C2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value)) for c in col) + 2
    wb.save(caminho)


def salvar_json(caminho, cabecalho, linhas, disc):
    dias = cabecalho[2:]
    dados = {
        "id": disc["id"],
        "name": disc["name"],
        "professor": disc["professor"],
        "horarios": disc["horarios"],
        "total_alunos": disc["total_alunos"],
        "alunos_presentes": disc["alunos_presentes"],
        "alunos": [
            {"nome": nome, "matricula": mat, "presenca": dict(zip(dias, status))}
            for nome, mat, *status in linhas
        ],
    }
    caminho.write_text(json.dumps(dados, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def salvar_relatorio():
    with open(RELATORIO, "w", encoding="utf-8") as f:
        for secao, itens in relatorio.items():
            f.write(f"== {secao} ({len(itens)}) ==\n")
            f.writelines(f"  {i}\n" for i in itens)
            f.write("\n")


def main():
    por_nome, por_email = carregar_inscricoes()
    eventos = carregar_eventos()
    presencas = mapear_presencas(eventos, por_nome, por_email)

    meta = json.loads(METADATA.read_text(encoding="utf-8"))
    turmas = [(disc, ler_alunos(PASTA_DISCIPLINAS / disc["id"])) for disc in meta["disciplinas"]]
    corrigir_matriculas(presencas, por_nome, por_email, turmas)

    for disc, alunos in turmas:
        pasta = PASTA_DISCIPLINAS / disc["id"]
        cabecalho, linhas, presentes = montar_tabela(disc, alunos, presencas)
        disc["total_alunos"] = str(len(alunos))
        disc["alunos_presentes"] = str(presentes)

        salvar_csv(pasta / "presença.csv", cabecalho, linhas)
        salvar_xlsx(pasta / "presença.xlsx", cabecalho, linhas, disc)
        salvar_json(pasta / "presença.json", cabecalho, linhas, disc)
        relatorio["Resumo por disciplina"].append(
            f"{disc['id']} - {disc['name']}: {presentes}/{len(alunos)} alunos com ao menos uma presença")
        print(f"{disc['id']}: {presentes}/{len(alunos)} presentes")

    METADATA.write_text(json.dumps(meta, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    salvar_relatorio()
    print(f"Relatório salvo em {RELATORIO.name}")


if __name__ == "__main__":
    main()
